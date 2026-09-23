from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ShelfSense_Project_Documentation.docx"
ACCENT = "3E6250"
LIGHT_GREEN = "EAF1EC"
LIGHT_GRAY = "F4F5F2"
BORDER = "D9DED9"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_run_font(run, name="Aptos", size=None, color="243029", bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("243029")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Aptos Display" if name in {"Title", "Heading 1"} else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ShelfSense project documentation")
    set_run_font(run, size=8, color="68746C")


def add_title(document, title, subtitle):
    paragraph = document.add_paragraph(style="Title")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    set_run_font(run, name="Aptos Display", size=28, color="000000", bold=True)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(subtitle)
    set_run_font(run, size=12, color=ACCENT, bold=True)


def add_paragraph(document, text, bold_lead=None):
    paragraph = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, ACCENT)
        set_cell_border(cell)
        set_cell_padding(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9, color="FFFFFF", bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            set_cell_shading(cells[index], "FFFFFF" if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_border(cells[index])
            set_cell_padding(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[index].paragraphs[0].runs:
                set_run_font(run, size=9)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_image(document, path, width=6.2, caption=None):
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(caption)
        set_run_font(run, size=8, color="68746C", italic=True)


def build_document():
    document = Document()
    style_document(document)
    add_title(document, "ShelfSense Project Documentation", "Food waste prevention and dynamic pricing dashboard")
    add_paragraph(
        document,
        "Dokumen ini menjelaskan apa yang dibangun, bagaimana cara menjalankannya, bagaimana membaca hasil analisis, dan kapan file Excel digunakan. ShelfSense adalah aplikasi portfolio berbasis Streamlit yang membantu tim bakery memahami risiko waste dan menentukan tindakan operasional yang dapat dijelaskan.",
    )
    add_image(document, ROOT / "assets" / "shelfsense-hero.png", width=6.1, caption="Tampilan visual utama ShelfSense")

    document.add_heading("Ringkasan project", level=1)
    add_paragraph(document, "ShelfSense menghubungkan produksi, penjualan, stok, shelf life, expiry date, dan waste dalam satu dashboard. Fokusnya bukan membuat prediksi yang terlihat canggih, melainkan membuat keputusan operasional lebih mudah dilacak.")
    add_table(
        document,
        ["Bagian", "Penjelasan"],
        [
            ("Masalah", "Produksi berlebih membuat makanan terbuang, sedangkan produksi terlalu sedikit dapat menyebabkan stockout."),
            ("Keputusan", "Produk mana yang perlu markdown, kapan produksi perlu dikurangi, dan cabang mana yang perlu diperhatikan."),
            ("Data", "Dataset simulasi 8 minggu, 3 cabang bakery, dan 12 produk."),
            ("Output", "KPI, chart, recommendation engine, chatbot terstruktur, database SQLite, dan workbook Excel."),
        ],
        widths=[1.2, 5.4],
    )

    document.add_heading("Cara menjalankan aplikasi", level=1)
    add_paragraph(document, "Buka terminal pada folder project, lalu jalankan perintah berikut:")
    code = document.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.space_after = Pt(10)
    run = code.add_run("python -m streamlit run app.py")
    set_run_font(run, name="Consolas", size=10, color=ACCENT, bold=True)
    add_paragraph(document, "Setelah server berjalan, buka alamat lokal yang ditampilkan di terminal. Saat pertama kali dibuka, aplikasi menampilkan popup pengantar. Pilih Open dashboard untuk masuk ke halaman utama.")

    document.add_heading("Cara memakai dashboard", level=1)
    add_table(
        document,
        ["Halaman", "Kapan digunakan", "Yang perlu diperhatikan"],
        [
            ("Today's Operations", "Melihat kondisi operasional saat ini.", "KPI, perubahan waste rate, dan daftar produk yang perlu perhatian."),
            ("Waste Monitor", "Mencari sumber waste terbesar.", "Bandingkan produk, alasan waste, dan cabang."),
            ("Markdown Planner", "Memeriksa rekomendasi markdown.", "Catat keputusan manager dan status pelaksanaannya."),
            ("Production Review", "Membandingkan produksi dengan demand baseline.", "Forecast adalah moving average sederhana untuk demo."),
            ("Import Data", "Memasukkan transaksi baru dari Excel.", "Baris yang tidak valid akan ditolak sebelum masuk database."),
            ("Ask ShelfSense", "Mendapatkan jawaban cepat dari topik yang didukung.", "Pertanyaan harus berkaitan dengan waste, cabang, loss, markdown, atau driver produk."),
            ("Data Notes", "Memeriksa definisi metrik dan mengunduh workbook.", "Semua impact diberi label estimated."),
        ],
        widths=[1.35, 2.35, 2.9],
    )
    add_paragraph(document, "Gunakan filter Branch untuk membatasi analisis ke satu cabang. Gunakan As of untuk mengubah tanggal acuan expiry dan rekomendasi.")

    document.add_heading("Fungsi upload Excel", level=1)
    add_paragraph(document, "Upload Excel digunakan untuk memasukkan transaksi baru ke database lokal. Fitur ini penting ketika data demo diganti dengan data operasional baru atau ketika ingin menguji bagaimana dashboard merespons batch transaksi tambahan.")
    add_table(
        document,
        ["Tahap", "Yang terjadi"],
        [
            ("Upload", "User memilih file .xlsx atau .xls dari halaman Import Data."),
            ("Normalisasi", "Nama kolom dirapikan agar variasi penulisan tetap dapat dibaca sistem."),
            ("Validasi", "Sistem memeriksa kolom wajib, angka, tanggal, duplikasi, quantity negatif, dan batas produksi."),
            ("Insert", "Hanya baris valid yang ditulis ke data/shelfsense.db."),
            ("Refresh", "KPI, chart, rekomendasi, dan jawaban chatbot memakai data terbaru."),
        ],
        widths=[1.2, 5.4],
    )
    add_paragraph(document, "Kolom minimum meliputi transaction_id, transaction_date, transaction_time, branch, product_id, product_name, category, production_qty, sold_qty, waste_qty, stock_qty, unit_price, unit_cost, discount_pct, expiry_date, dan shelf_life_days.")

    document.add_heading("Mengapa hasilnya bisa seperti itu", level=1)
    add_paragraph(document, "Data yang dipakai adalah simulasi yang sengaja memiliki variasi operasional. Karena itu, hasil dashboard tidak dibuat selalu positif. Ada produk yang lambat terjual, cabang dengan waste lebih tinggi, dan hari tertentu yang menghasilkan estimated loss lebih besar.")
    add_bullets(
        document,
        [
            "Weekend memiliki demand lebih tinggi.",
            "Hujan menurunkan demand pada beberapa produk.",
            "Beberapa produk mengalami overproduction untuk menunjukkan dampak stok berlebih.",
            "Produk dengan shelf life pendek lebih mudah masuk markdown window.",
            "Produk yang tetap lambat terjual setelah diskon ditandai sebagai slow moving.",
        ],
    )
    add_paragraph(document, "Hasil ini bukan klaim performa bisnis nyata. Nilainya menunjukkan bagaimana aturan analitik bekerja pada dataset demo, sehingga pembaca dapat mengikuti hubungan antara data, rumus, dan rekomendasi.")

    document.add_heading("Definisi metrik dan aturan rekomendasi", level=1)
    add_table(
        document,
        ["Metrik atau aturan", "Definisi"],
        [
            ("Sold-through rate", "sold_qty dibagi production_qty."),
            ("Waste rate", "waste_qty dibagi production_qty."),
            ("Estimated loss", "waste_qty dikali unit_cost."),
            ("Markdown 20 persen", "Expiry tersisa maksimal 1 hari dan stok melebihi demand baseline."),
            ("Markdown 10 persen", "Expiry tersisa 2 hari dan stok masih surplus."),
            ("Reduce production", "Waste rate melewati 10 persen selama pola operasional menunjukkan masalah berulang."),
            ("Demand baseline", "Moving average sederhana dari penjualan sebelumnya."),
        ],
        widths=[1.8, 4.8],
    )
    add_paragraph(document, "Estimated revenue dan estimated loss adalah perhitungan indikatif. Nilai aktual setelah manager mengambil tindakan perlu dimasukkan sebagai hasil operasional terpisah.")

    document.add_heading("Struktur folder", level=1)
    add_table(
        document,
        ["Lokasi", "Peran"],
        [
            ("app.py", "Entry point Streamlit dan alur startup aplikasi."),
            ("src/dashboard.py", "Renderer halaman dan alur interaksi dashboard."),
            ("src/ui.py", "Style, sidebar, hero section, popup, dan helper tampilan."),
            ("src/analytics.py", "Perhitungan KPI, trend, forecast baseline, dan risk frame."),
            ("src/decision_engine.py", "Aturan rekomendasi markdown dan pengurangan produksi."),
            ("src/data_loader.py", "Pembacaan dan validasi file Excel."),
            ("src/database.py", "Schema dan operasi SQLite."),
            ("tests/", "Unit test dan integration test."),
            ("data/", "Dataset simulasi dan database lokal."),
            ("exports/", "Workbook yang bisa dibuka manual di Excel."),
            ("assets/", "Gambar hero, product board, dan chart PNG."),
            ("docs/", "Project map, user guide, spec, plan, dan dokumen ini."),
        ],
        widths=[1.9, 4.7],
    )

    document.add_heading("File data yang bisa dibuka manual", level=1)
    add_bullets(
        document,
        [
            "data/sample_food_waste.xlsx adalah dataset simulasi mentah.",
            "data/shelfsense.db adalah database SQLite lokal.",
            "exports/ShelfSense_demo_workbook.xlsx adalah versi Excel dengan sheet Transactions, KPI Summary, Waste by Product, dan Recommendations.",
        ],
    )
    add_paragraph(document, "Untuk melihat data dengan cara yang paling mudah, buka ShelfSense_demo_workbook.xlsx menggunakan Excel atau aplikasi spreadsheet lain.")

    document.add_heading("Catatan pengembangan", level=1)
    add_paragraph(document, "Struktur aplikasi sengaja dibuat sederhana dan modular. app.py hanya mengurus startup, sedangkan logika tampilan dan logika analitik dipisahkan ke modul yang lebih kecil. Pendekatan ini membuat project lebih mudah dibaca dan lebih aman untuk dikembangkan tanpa mengubah seluruh aplikasi.")
    add_paragraph(document, "MVP ini menggunakan SQLite dan moving average sederhana. Untuk penggunaan bisnis sungguhan, tahap berikutnya adalah menambahkan autentikasi, database produksi, forecasting yang divalidasi, audit log keputusan, dan deployment dengan secret management.")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
