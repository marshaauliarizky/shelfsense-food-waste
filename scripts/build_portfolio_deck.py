from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ShelfSense_Portfolio_Deck_v2.pptx"
ASSETS = ROOT / "assets"

INK = RGBColor(36, 48, 41)
MUTED = RGBColor(104, 116, 108)
PAPER = RGBColor(247, 244, 237)
GREEN = RGBColor(88, 117, 104)
AMBER = RGBColor(196, 134, 56)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(224, 224, 215)
DARK = RGBColor(37, 51, 44)


def set_background(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(0.02)
    box.text_frame.margin_right = Inches(0.02)
    box.text_frame.margin_top = Inches(0.02)
    box.text_frame.margin_bottom = Inches(0.02)
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_box(slide, x, y, w, h, fill=WHITE, line=LINE, rounded=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def add_image(slide, path, x, y, w, h=None):
    if not path.exists():
        return
    if h is None:
        h = w * 0.6
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_footer(slide, number):
    add_text(slide, "ShelfSense  |  Portfolio project", 0.6, 7.12, 4.8, 0.16, size=8, color=MUTED)
    add_text(slide, f"{number:02d}", 11.5, 7.12, 0.5, 0.16, size=8, color=MUTED, align=PP_ALIGN.CENTER)


def add_title(slide, heading, subheading=None):
    add_text(slide, heading, 0.65, 0.42, 10.4, 0.42, size=28, bold=True)
    if subheading:
        add_text(slide, subheading, 0.67, 0.96, 10.6, 0.26, size=13, color=MUTED)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_image(slide, ASSETS / "shelfsense-hero.png", 6.45, 0.0, 3.55, 5.65)
    dark_panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.0), Inches(0), Inches(3.333), Inches(7.5))
    dark_panel.fill.solid(); dark_panel.fill.fore_color.rgb = DARK; dark_panel.line.fill.background()
    add_text(slide, "SHELF\nSENSE", 10.55, 1.05, 2.1, 1.15, size=30, color=WHITE, bold=True)
    add_text(slide, "Food waste prevention\nand markdown planning", 10.55, 2.55, 2.2, 0.7, size=16, color=RGBColor(213, 226, 215))
    add_text(slide, "Portfolio case study", 0.72, 0.9, 4.0, 0.22, size=12, color=AMBER, bold=True)
    add_text(slide, "What should sell,\nwhat may spoil,\nand what to do next", 0.72, 1.55, 5.4, 2.1, size=34, bold=True)
    add_text(slide, "A practical Streamlit dashboard for bakery operations", 0.75, 4.1, 4.6, 0.45, size=16, color=MUTED)
    add_text(slide, "Built with Python, Pandas, SQLite, Plotly, and Streamlit", 0.75, 6.35, 5.4, 0.22, size=11, color=MUTED)

    slide = prs.slides.add_slide(blank); set_background(slide)
    add_title(slide, "The operating problem", "Perishable inventory creates a trade off between availability and waste")
    add_text(slide, "Bakery teams need to make decisions before products expire. A useful tool has to show the evidence behind the decision, not only a final score.", 0.72, 1.6, 5.2, 1.05, size=20)
    add_box(slide, 0.72, 3.05, 2.65, 1.55, fill=RGBColor(235, 242, 236), line=RGBColor(203, 220, 204))
    add_text(slide, "Too much production", 0.98, 3.35, 2.1, 0.3, size=18, bold=True)
    add_text(slide, "Creates surplus stock\nand higher waste risk", 0.98, 3.85, 1.95, 0.5, size=14, color=MUTED)
    add_box(slide, 3.72, 3.05, 2.65, 1.55, fill=RGBColor(255, 250, 240), line=RGBColor(231, 216, 188))
    add_text(slide, "Too little production", 3.98, 3.35, 2.1, 0.3, size=18, bold=True)
    add_text(slide, "Increases stockout risk\nand missed sales", 3.98, 3.85, 1.95, 0.5, size=14, color=MUTED)
    add_image(slide, ASSETS / "bakery-product-board.png", 7.35, 1.65, 4.55, 3.1)
    add_text(slide, "ShelfSense keeps the decision visible: production, sales, stock, expiry, waste, and markdown are read together.", 7.35, 5.15, 4.55, 0.65, size=15, color=MUTED)
    add_footer(slide, 2)

    slide = prs.slides.add_slide(blank); set_background(slide)
    add_title(slide, "The dashboard in use", "A manager starts with the current operational picture")
    add_image(slide, ASSETS / "screenshots" / "01_operations.png", 0.55, 1.45, 7.1, 5.45)
    add_text(slide, "The home screen answers three questions", 8.05, 1.75, 4.2, 0.35, size=20, bold=True)
    add_text(slide, "1  How much was produced and sold?\n\n2  Where is waste moving?\n\n3  Which product needs attention before expiry?", 8.05, 2.35, 4.0, 1.8, size=16)
    add_text(slide, "The Branch and As of filters keep the view grounded in a specific operating scope.", 8.05, 5.2, 3.9, 0.7, size=15, color=MUTED)
    add_footer(slide, 3)

    slide = prs.slides.add_slide(blank); set_background(slide)
    add_title(slide, "The analysis behind the screen", "Simple formulas keep the recommendations explainable")
    metrics = [
        ("Sold-through rate", "sold quantity / production quantity"),
        ("Waste rate", "waste quantity / production quantity"),
        ("Estimated loss", "waste quantity × unit cost"),
        ("Stock gap", "current stock − demand baseline"),
    ]
    y = 1.72
    for label, description in metrics:
        add_text(slide, label, 0.9, y, 2.4, 0.28, size=18, bold=True)
        add_text(slide, description, 3.6, y + 0.02, 3.6, 0.26, size=16, color=MUTED)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(y + 0.55), Inches(6.6), Inches(0.01))
        line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
        y += 1.0
    add_image(slide, ASSETS / "chart-sales-waste.png", 8.55, 1.68, 3.7, 2.35)
    add_image(slide, ASSETS / "chart-waste-reasons.png", 8.55, 4.25, 3.7, 2.35)
    add_footer(slide, 4)

    slide = prs.slides.add_slide(blank); set_background(slide)
    add_title(slide, "From transaction to recommendation", "The pipeline keeps data, rules, and presentation separate")
    steps = [
        ("01", "Excel", "Incoming transactions"),
        ("02", "Validation", "Clean rows only"),
        ("03", "SQLite", "Local source of truth"),
        ("04", "Analytics", "KPI and demand baseline"),
        ("05", "Decision engine", "Markdown and production rules"),
        ("06", "Dashboard", "Manager view and notes"),
    ]
    x = 0.65
    for number, label, description in steps:
        add_box(slide, x, 2.55, 1.78, 1.6, fill=WHITE, line=LINE)
        add_text(slide, number, x + 0.18, 2.76, 0.45, 0.22, size=12, color=AMBER, bold=True)
        add_text(slide, label, x + 0.18, 3.18, 1.4, 0.28, size=17, bold=True)
        add_text(slide, description, x + 0.18, 3.68, 1.4, 0.35, size=12, color=MUTED)
        x += 2.08
    add_text(slide, "This separation makes the project easier to test and extend. A rule can change without rewriting the interface, and the interface can evolve without changing the database schema.", 1.0, 5.35, 11.2, 0.7, size=17, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 5)

    slide = prs.slides.add_slide(blank); set_background(slide)
    add_title(slide, "Excel upload and auditability", "The upload flow is part of the product, not an afterthought")
    add_text(slide, "A new workbook is checked before it can change the dashboard.", 0.72, 1.55, 5.8, 0.4, size=20, bold=True)
    add_text(slide, "The system checks required columns, dates, numeric values, negative quantities, duplicate transaction IDs, discount ranges, and production limits.", 0.72, 2.2, 5.7, 1.1, size=17, color=MUTED)
    add_box(slide, 0.72, 4.0, 5.45, 1.1, fill=RGBColor(235, 242, 236), line=RGBColor(203, 220, 204))
    add_text(slide, "Only valid rows are inserted", 1.0, 4.28, 3.3, 0.28, size=19, bold=True)
    add_text(slide, "The dashboard then refreshes from SQLite.", 1.0, 4.68, 3.8, 0.25, size=15, color=MUTED)
    add_image(slide, ASSETS / "screenshots" / "import_data.png", 7.2, 1.45, 5.25, 5.05)
    add_footer(slide, 6)

    slide = prs.slides.add_slide(blank); set_background(slide, DARK)
    add_text(slide, "What this project demonstrates", 0.75, 0.72, 8.2, 0.55, size=30, color=WHITE, bold=True)
    add_text(slide, "A portfolio project should show judgment as well as implementation.", 0.78, 1.38, 8.4, 0.3, size=16, color=RGBColor(213, 226, 215))
    add_text(slide, "DATA\n\nSimulated data with clear assumptions and validation", 1.0, 2.55, 3.0, 1.4, size=19, color=WHITE, bold=True)
    add_text(slide, "DECISIONS\n\nExplainable rules connect metrics to manager actions", 5.0, 2.55, 3.1, 1.4, size=19, color=WHITE, bold=True)
    add_text(slide, "DELIVERY\n\nA usable dashboard, exportable workbook, tests, and documentation", 9.0, 2.55, 3.1, 1.55, size=19, color=WHITE, bold=True)
    add_text(slide, "Current scope: local MVP with SQLite and a moving average baseline. The next production steps are authenticated access, a managed database, validated forecasting, and deployment.", 1.0, 5.45, 11.2, 0.72, size=16, color=RGBColor(213, 226, 215), align=PP_ALIGN.CENTER)
    add_footer(slide, 7)

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build_deck()
